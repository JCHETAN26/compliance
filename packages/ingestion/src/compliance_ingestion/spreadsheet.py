"""XLSX and CSV extraction into structured rows plus one searchable text line per row.

Header detection: among the first ``HEADER_WINDOW`` non-empty rows, the header is the
first row whose non-empty cells are all text, distinct, at least two, and cover at least
60% of the sheet's used width. If no row qualifies, columns are named ``Column A``,
``Column B``... and every row is data.

Canonical text per sheet::

    Sheet: Vendors
    Vendor name: Acme Ltd; Data access level: High; Contract file name: acme_nda.pdf

Each data row's line is one chunk (offsets point at that line).
"""

from __future__ import annotations

import csv
import io
import math
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from compliance_ingestion.textbuilder import BLOCK_SEPARATOR, LINE_SEPARATOR, TextBuilder
from compliance_ingestion.types import CellValue, ExtractionError, RowInfo, SheetInfo

HEADER_WINDOW = 10
MAX_COLUMNS = 200
MAX_CELL_CHARS = 2000


@dataclass
class SheetData:
    name: str
    rows: list[tuple[int, list[Any]]]  # (1-based source row number, values)
    truncated: bool = False


@dataclass
class SpreadsheetExtraction:
    builder: TextBuilder
    sheets: list[SheetInfo]
    rows: list[RowInfo]
    warnings: list[str] = field(default_factory=list)


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(timespec="minutes")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return str(value)
        return str(int(value)) if value.is_integer() else f"{value:.15g}"
    if isinstance(value, Decimal):
        return format(value.normalize(), "f")
    text = str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ").strip()
    return text[:MAX_CELL_CHARS]


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, bool | int | str):
        return value if not isinstance(value, str) else value.strip()[:MAX_CELL_CHARS]
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, datetime | date | time):
        return display_value(value)
    return display_value(value)


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _looks_numeric(text: str) -> bool:
    try:
        float(text.replace(",", ""))
    except ValueError:
        return False
    return True


def detect_header(rows: list[list[Any]], width: int) -> int | None:
    """Index into ``rows`` of the header row, or None."""
    checked = 0
    for index, row in enumerate(rows):
        values = [v for v in row if not _is_empty(v)]
        if not values:
            continue
        checked += 1
        if checked > HEADER_WINDOW:
            break
        texts = [str(v).strip() for v in values if isinstance(v, str)]
        if (
            len(values) >= 2
            and len(texts) == len(values)
            and not any(_looks_numeric(t) for t in texts)
            and len({t.lower() for t in texts}) == len(texts)
            and len(values) >= 0.6 * width
        ):
            return index
    return None


def _column_names(header: list[Any] | None, width: int) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        raw = header[i] if header is not None and i < len(header) else None
        name = display_value(raw) if not _is_empty(raw) else f"Column {get_column_letter(i + 1)}"
        name = name[:120]
        key = name.lower()
        if key in seen:
            seen[key] += 1
            name = f"{name} ({seen[key]})"
        else:
            seen[key] = 1
        names.append(name)
    return names


def _build(sheets: list[SheetData], warnings: list[str]) -> SpreadsheetExtraction:
    builder = TextBuilder()
    infos: list[SheetInfo] = []
    rows_out: list[RowInfo] = []
    for sheet_index, sheet in enumerate(sheets):
        values = [row for _, row in sheet.rows]
        width = 0
        for row in values:
            for i in range(min(len(row), MAX_COLUMNS) - 1, -1, -1):
                if not _is_empty(row[i]):
                    width = max(width, i + 1)
                    break
        if width == 0:
            infos.append(
                SheetInfo(
                    sheet_index=sheet_index,
                    sheet_name=sheet.name,
                    columns=[],
                    header_row_number=None,
                    row_count=0,
                    truncated=sheet.truncated,
                )
            )
            continue
        header_idx = detect_header(values, width)
        columns = _column_names(values[header_idx] if header_idx is not None else None, width)
        builder.separate(BLOCK_SEPARATOR)
        builder.append(f"Sheet: {sheet.name}")
        count = 0
        start_at = header_idx + 1 if header_idx is not None else 0
        for row_number, row in sheet.rows[start_at:]:
            cells = [
                CellValue(
                    column=columns[i], value=json_value(row[i]), display=display_value(row[i])
                )
                for i in range(min(width, len(row)))
            ]
            if all(not c.display for c in cells):
                continue
            line = "; ".join(f"{c.column}: {c.display}" for c in cells if c.display)
            builder.append(LINE_SEPARATOR)
            start = builder.length
            builder.append(line)
            rows_out.append(
                RowInfo(
                    sheet_index=sheet_index,
                    row_number=row_number,
                    cells=cells,
                    char_start=start,
                    char_end=builder.length,
                )
            )
            count += 1
        infos.append(
            SheetInfo(
                sheet_index=sheet_index,
                sheet_name=sheet.name,
                columns=columns,
                header_row_number=sheet.rows[header_idx][0] if header_idx is not None else None,
                row_count=count,
                truncated=sheet.truncated,
            )
        )
        if sheet.truncated:
            warnings.append(f"Sheet {sheet.name!r} was truncated at the row limit.")
    if not rows_out:
        warnings.append("The spreadsheet contains no data rows.")
    return SpreadsheetExtraction(builder=builder, sheets=infos, rows=rows_out, warnings=warnings)


def extract_xlsx(path: Path, max_rows: int) -> SpreadsheetExtraction:
    # A file handle, not a path: stored originals have no extension, and openpyxl refuses
    # paths that do not end in .xlsx.
    with path.open("rb") as handle:
        try:
            # read_only streams rows; data_only returns cached formula results, never formulas.
            workbook = load_workbook(handle, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for corrupt input
            raise ExtractionError("The XLSX file is damaged or not a valid workbook.") from exc
        sheets: list[SheetData] = []
        try:
            for worksheet in workbook.worksheets:
                rows: list[tuple[int, list[Any]]] = []
                truncated = False
                for number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                    if len(rows) >= max_rows + HEADER_WINDOW:
                        truncated = True
                        break
                    rows.append((number, list(row[:MAX_COLUMNS])))
                sheets.append(SheetData(name=str(worksheet.title), rows=rows, truncated=truncated))
        finally:
            workbook.close()
    return _build(sheets, [])


def decode_csv(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def extract_csv(path: Path, max_rows: int, sheet_name: str = "CSV") -> SpreadsheetExtraction:
    text, encoding = decode_csv(path.read_bytes())
    if "\x00" in text:
        raise ExtractionError("The CSV file contains binary data.")
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:65536], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    rows: list[tuple[int, list[Any]]] = []
    truncated = False
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    try:
        for number, row in enumerate(reader, start=1):
            if len(rows) >= max_rows + HEADER_WINDOW:
                truncated = True
                break
            rows.append((number, list(row[:MAX_COLUMNS])))
    except csv.Error as exc:
        raise ExtractionError(f"The CSV file could not be parsed: {exc}") from exc
    warnings = [] if encoding == "utf-8-sig" else [f"CSV decoded as {encoding} (not valid UTF-8)."]
    return _build([SheetData(name=sheet_name, rows=rows, truncated=truncated)], warnings)
