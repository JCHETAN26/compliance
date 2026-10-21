# mypy: disallow-untyped-calls=False
# (fixtures are built with PyMuPDF/python-docx, which ship partial annotations)
"""Parser tests: PDF text layer (exact boxes), OCR fallback, DOCX, XLSX and CSV."""

from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

import docx
import pymupdf
import pytest
from openpyxl import Workbook
from PIL import Image

from compliance_ingestion.chunking import ChunkOptions
from compliance_ingestion.docx_parser import extract_docx
from compliance_ingestion.layout import merge_char_boxes
from compliance_ingestion.pdf import PdfOptions, extract_pdf
from compliance_ingestion.pipeline import PipelineOptions, extract_document
from compliance_ingestion.spreadsheet import detect_header, extract_csv, extract_xlsx
from compliance_ingestion.types import ExtractionError

PDF_OPTIONS = PdfOptions(
    max_pages=50, ocr_language="eng", ocr_dpi=300, ocr_min_text_chars=25, ocr_timeout=120
)


class WordCounter:
    def count_many(self, texts: list[str]) -> list[int]:
        return [len(t.split()) for t in texts]


def pipeline_options(min_tokens: int = 200) -> PipelineOptions:
    return PipelineOptions(
        pdf=PDF_OPTIONS, chunks=ChunkOptions(min_tokens, 400, 40), max_sheet_rows=1000
    )


# --------------------------------------------------------------------------- PDF text


def _text_pdf(path: Path, lines: list[tuple[float, float, str, float]], rotate: int = 0) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    for x, y, text, size in lines:
        page.insert_text((x, y), text, fontsize=size, fontname="helv")
    if rotate:
        page.set_rotation(rotate)
    doc.save(path)
    doc.close()
    return path


def test_pdf_text_offsets_and_exact_boxes(tmp_path: Path) -> None:
    path = _text_pdf(
        tmp_path / "a.pdf",
        [
            (72, 100, "1. Confidential Information", 14),
            (72, 130, "The Recipient shall keep all information secret.", 11),
            (72, 145, "It shall not disclose it to any third party.", 11),
        ],
    )
    result = extract_pdf(path, PDF_OPTIONS)
    text = result.builder.text()
    assert "The Recipient shall keep all information secret." in text
    assert len(result.builder.boxes) == len(text)
    with pymupdf.open(path) as doc:
        page = doc[0]
        for needle in ("Recipient shall keep", "third party"):
            start = text.index(needle)
            rects = merge_char_boxes(result.builder.boxes, start, start + len(needle))
            assert len(rects) == 1
            r = rects[0]
            # Independent check: PyMuPDF reads back exactly the cited words from the
            # highlighted region. Glyph boxes of tightly spaced lines touch, so the check
            # uses the rect's central band (full width) to exclude adjacent lines.
            band = (r.y1 - r.y0) * 0.3
            probe = pymupdf.Rect(r.x0, r.y0 + band, r.x1, r.y1 - band)
            assert page.get_textbox(probe).strip() == needle
    page = result.pages[0]
    assert (page.width, page.height, page.is_ocr) == (612, 792, False)
    assert text[page.char_start : page.char_end].startswith("1. Confidential Information")


def test_multiline_span_produces_one_rect_per_line(tmp_path: Path) -> None:
    path = _text_pdf(
        tmp_path / "b.pdf",
        [(72, 200, "First line of the clause", 11), (72, 215, "continues on the second line.", 11)],
    )
    result = extract_pdf(path, PDF_OPTIONS)
    text = result.builder.text()
    start = text.index("First line")
    end = text.index("second line.") + len("second line.")
    rects = merge_char_boxes(result.builder.boxes, start, end)
    assert len(rects) == 2
    assert rects[0].y1 <= rects[1].y0 + 2
    assert all(abs(r.x0 - 72) < 1.0 for r in rects)


def test_rotated_page_boxes_are_in_displayed_orientation(tmp_path: Path) -> None:
    path = _text_pdf(
        tmp_path / "r.pdf", [(72, 100, "Rotated confidentiality clause", 12)], rotate=90
    )
    result = extract_pdf(path, PDF_OPTIONS)
    page = result.pages[0]
    assert (page.width, page.height) == (792, 612)
    text = result.builder.text()
    start = text.index("Rotated")
    rects = merge_char_boxes(
        result.builder.boxes, start, start + len("Rotated confidentiality clause")
    )
    assert rects
    for r in rects:
        assert 0 <= r.x0 < r.x1 <= page.width
        assert 0 <= r.y0 < r.y1 <= page.height
    # Horizontal text on a page rotated 90 degrees is displayed vertically.
    union_h = max(r.y1 for r in rects) - min(r.y0 for r in rects)
    union_w = max(r.x1 for r in rects) - min(r.x0 for r in rects)
    assert union_h > union_w


def test_heading_detection_and_sections(tmp_path: Path) -> None:
    path = _text_pdf(
        tmp_path / "h.pdf",
        [
            (72, 80, "MUTUAL NON-DISCLOSURE AGREEMENT", 16),
            (72, 120, "2. Obligations", 11),
            (72, 150, "2.1 The Recipient shall protect the information.", 11),
        ],
    )
    result = extract_document(path, "pdf", pipeline_options(), WordCounter())
    kinds = {result.text[b.char_start : b.char_end]: b.kind for b in result.blocks}
    assert kinds["MUTUAL NON-DISCLOSURE AGREEMENT"] == "heading"
    clause = next(b for b in result.blocks if result.text[b.char_start :].startswith("2.1"))
    # A numbered paragraph whose first sentence is not a short title is labelled by number.
    assert clause.section_path == "2. Obligations > 2.1"
    assert result.doc_type == "contract"


def test_encrypted_damaged_and_page_limited_pdfs_rejected(tmp_path: Path) -> None:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    aes256 = pymupdf.PDF_ENCRYPT_AES_256  # type: ignore[attr-defined]
    doc.save(tmp_path / "enc.pdf", encryption=aes256, user_pw="pw", owner_pw="pw2")
    with pytest.raises(ExtractionError, match="Password-protected"):
        extract_pdf(tmp_path / "enc.pdf", PDF_OPTIONS)
    (tmp_path / "bad.pdf").write_bytes(b"%PDF-1.7\n garbage without objects")
    with pytest.raises(ExtractionError):
        extract_pdf(tmp_path / "bad.pdf", PDF_OPTIONS)
    many = pymupdf.open()
    for _ in range(3):
        many.new_page().insert_text((72, 72), "page text that is long enough to count")
    many.save(tmp_path / "many.pdf")
    with pytest.raises(ExtractionError, match="limit is 2"):
        extract_pdf(tmp_path / "many.pdf", PdfOptions(2, "eng", 300, 25, 60))


# --------------------------------------------------------------------------- OCR


def _scanned_pdf(path: Path, lines: list[tuple[float, float, str]]) -> Path:
    """Render text to an image and embed only the image: a genuine image-only PDF."""
    source = pymupdf.open()
    page = source.new_page(width=612, height=792)
    for x, y, text in lines:
        page.insert_text((x, y), text, fontsize=16, fontname="helv")
    pix = page.get_pixmap(dpi=200, colorspace=pymupdf.csGRAY)
    image = Image.frombytes("L", (pix.width, pix.height), pix.samples)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    scan = pymupdf.open()
    target = scan.new_page(width=612, height=792)
    target.insert_image(target.rect, stream=buffer.getvalue())
    scan.save(path)
    return path


def test_image_only_page_is_ocrd_with_page_coordinates(tmp_path: Path) -> None:
    path = _scanned_pdf(
        tmp_path / "scan.pdf",
        [
            (72, 200, "CONFIDENTIAL INFORMATION SHALL BE RETURNED"),
            (72, 400, "Governing law is England."),
        ],
    )
    with pymupdf.open(path) as doc:
        assert doc[0].get_text().strip() == ""  # no text layer at all
    result = extract_pdf(path, PDF_OPTIONS)
    page = result.pages[0]
    assert page.is_ocr
    assert page.ocr_confidence is not None
    assert page.ocr_confidence > 60
    text = result.builder.text()
    assert "CONFIDENTIAL INFORMATION SHALL BE RETURNED" in text
    start = text.index("RETURNED")
    rects = merge_char_boxes(result.builder.boxes, start, start + len("RETURNED"))
    assert len(rects) == 1
    r = rects[0]
    # Text was drawn with its baseline at y=200 (16pt): the OCR box must sit there.
    assert 180 <= r.y0 <= 200 <= r.y1 <= 210
    assert 72 < r.x0 < r.x1 < 612
    gov = text.index("Governing")
    gov_rect = merge_char_boxes(result.builder.boxes, gov, gov + 9)[0]
    assert 380 <= gov_rect.y0 <= 400


def test_blank_page_ocr_reports_warning(tmp_path: Path) -> None:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "A first page with plenty of real text on it.")
    doc.new_page()
    doc.save(tmp_path / "blank.pdf")
    result = extract_pdf(tmp_path / "blank.pdf", PDF_OPTIONS)
    assert result.pages[1].is_ocr
    assert result.pages[1].char_start == result.pages[1].char_end
    assert any("Page 2" in w for w in result.warnings)


# --------------------------------------------------------------------------- DOCX


def _docx(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("Non-Disclosure Agreement", level=0)
    document.add_heading("1. Definitions", level=1)
    document.add_paragraph(
        "Confidential Information means all data disclosed. It includes trade secrets."
    )
    document.add_paragraph("Do not copy.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Party"
    table.cell(0, 1).text = "Role"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "Acme Ltd"
    document.save(str(path))
    return path


def test_docx_paragraphs_styles_offsets(tmp_path: Path) -> None:
    result = extract_document(
        _docx(tmp_path / "a.docx"), "docx", pipeline_options(20), WordCounter()
    )
    by_text = {result.text[b.char_start : b.char_end]: b for b in result.blocks}
    assert by_text["Non-Disclosure Agreement"].kind == "heading"
    assert by_text["Non-Disclosure Agreement"].style == "Title"
    assert by_text["1. Definitions"].kind == "heading"
    assert by_text["Do not copy."].kind == "list_item"
    assert by_text["Party | Role"].kind == "table_row"
    assert by_text["Acme Ltd"].kind == "table_row", "merged cells appear once"
    para = by_text["Confidential Information means all data disclosed. It includes trade secrets."]
    assert para.section_path == "1. Definitions"
    assert result.pages == []
    for chunk in result.chunks:
        assert result.text[chunk.char_start : chunk.char_end] == chunk.text
        assert chunk.bboxes == []


def test_docx_corrupt_rejected(tmp_path: Path) -> None:
    (tmp_path / "x.docx").write_bytes(b"PK\x03\x04not really")
    with pytest.raises(ExtractionError):
        extract_docx(tmp_path / "x.docx")


# --------------------------------------------------------------------------- spreadsheets


def _register(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Vendors"
    sheet["A1"] = "Vendor register (title row)"
    sheet.append([])
    sheet.append(
        ["Vendor name", "Data access level", "Contract file name", "Renewal date", "Spend"]
    )
    sheet.append(["Acme Ltd", "High", "acme_nda.pdf", date(2027, 3, 31), 1200.0])
    sheet.append(["Beta GmbH", None, "beta.pdf", datetime(2026, 12, 1, 15, 30), 99.5])  # noqa: DTZ001
    sheet.append([None, None, None, None, None])
    sheet.append(["Gamma", "Low", None, None, 3])
    workbook.save(path)
    return path


def test_xlsx_header_after_title_rows_and_value_display(tmp_path: Path) -> None:
    result = extract_xlsx(_register(tmp_path / "r.xlsx"), max_rows=1000)
    sheet = result.sheets[0]
    assert sheet.header_row_number == 3
    assert sheet.columns == [
        "Vendor name",
        "Data access level",
        "Contract file name",
        "Renewal date",
        "Spend",
    ]
    assert sheet.row_count == 3, "blank rows are skipped"
    rows = {r.row_number: r for r in result.rows}
    assert set(rows) == {4, 5, 7}
    acme = {c.column: c for c in rows[4].cells}
    assert acme["Renewal date"].display == "2027-03-31"
    assert acme["Spend"].display == "1200"
    beta = {c.column: c for c in rows[5].cells}
    assert beta["Data access level"].display == ""
    assert beta["Data access level"].value is None
    assert beta["Renewal date"].display == "2026-12-01T15:30"
    text = result.builder.text()
    line = text[rows[4].char_start : rows[4].char_end]
    assert line == (
        "Vendor name: Acme Ltd; Data access level: High; Contract file name: acme_nda.pdf; "
        "Renewal date: 2027-03-31; Spend: 1200"
    )
    gamma_line = text[rows[7].char_start : rows[7].char_end]
    assert "Contract file name" not in gamma_line, "empty cells are omitted from row text"


def test_xlsx_rows_become_one_chunk_each(tmp_path: Path) -> None:
    result = extract_document(
        _register(tmp_path / "r.xlsx"), "xlsx", pipeline_options(), WordCounter()
    )
    assert result.doc_type == "spreadsheet"
    assert [c.row_number for c in result.chunks] == [4, 5, 7]
    for chunk, row in zip(result.chunks, result.rows, strict=True):
        assert chunk.text == result.text[row.char_start : row.char_end]
        assert chunk.section_path == f"Vendors > Row {row.row_number}"


def test_xlsx_truncation_reported(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Name", "Value"])
    for i in range(50):
        sheet.append([f"row {i}", i])
    workbook.save(tmp_path / "big.xlsx")
    result = extract_xlsx(tmp_path / "big.xlsx", max_rows=20)
    assert result.sheets[0].truncated
    assert any("truncated" in w for w in result.warnings)


def test_csv_dialect_encoding_and_header(tmp_path: Path) -> None:
    content = "Vendor;Access;Notes\nCafé Ltd;High;Größe ok\nBeta;Low;\n"
    (tmp_path / "a.csv").write_bytes(content.encode("cp1252"))
    result = extract_csv(tmp_path / "a.csv", max_rows=100)
    assert result.sheets[0].columns == ["Vendor", "Access", "Notes"]
    assert result.sheets[0].header_row_number == 1
    first = {c.column: c.display for c in result.rows[0].cells}
    assert first == {"Vendor": "Café Ltd", "Access": "High", "Notes": "Größe ok"}
    assert any("cp1252" in w for w in result.warnings)


def test_csv_without_header_gets_column_letters(tmp_path: Path) -> None:
    (tmp_path / "n.csv").write_text("1,2,3\n4,5,6\n", encoding="utf-8")
    result = extract_csv(tmp_path / "n.csv", max_rows=100)
    assert result.sheets[0].header_row_number is None
    assert result.sheets[0].columns == ["Column A", "Column B", "Column C"]
    assert len(result.rows) == 2


def test_header_detection_rejects_numeric_and_sparse_rows() -> None:
    assert detect_header([["1", "2"], ["Name", "Value"]], 2) == 1
    assert detect_header([["Only one"], ["Name", "Value"]], 2) == 1
    assert detect_header([[1, 2], [3, 4]], 2) is None


def test_xlsx_and_docx_parse_from_extensionless_storage_paths(tmp_path: Path) -> None:
    """Stored originals are named ``original`` (no extension)."""
    stored = tmp_path / "original"
    stored.write_bytes(_register(tmp_path / "r.xlsx").read_bytes())
    assert extract_xlsx(stored, max_rows=100).sheets[0].row_count == 3
    stored.write_bytes(_docx(tmp_path / "a.docx").read_bytes())
    assert extract_docx(stored).blocks
