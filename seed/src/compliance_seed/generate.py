"""Deterministic generation of the demo files (same inputs -> byte-identical outputs).

Determinism matters: uploads are de-duplicated by SHA-256, so re-running ``make seed``
must regenerate identical bytes to stay idempotent. PDFs are saved without a new file
ID and with fixed metadata dates. Office files (zip containers) are rewritten with fixed
entry timestamps and fixed core-property dates. Scan noise uses a seeded RNG.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import re
import zipfile
from dataclasses import dataclass
from importlib import resources
from pathlib import Path
from typing import Any, Literal

import docx
import numpy as np
import pymupdf
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image

FIXED_PDF_DATE = "D:20211005000000Z"
FIXED_DATETIME = dt.datetime(2021, 10, 5, 0, 0, 0)  # noqa: DTZ001 - zip/Office local timestamps
FIXED_ZIP_TIME = (2021, 10, 5, 0, 0, 0)
SCAN_DPI = 200
_CORE_DATE = re.compile(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)")

SourceFormat = Literal["original_pdf", "scanned_pdf", "docx"]


@dataclass(frozen=True)
class VendorSpec:
    name: str
    data_access: str
    renewal_date: dt.date
    contract_file: str
    contractnli_id: int | None
    source_format: SourceFormat | None


@dataclass(frozen=True)
class Manifest:
    dataset: dict[str, Any]
    demo: dict[str, str]
    vendors: list[VendorSpec]


@dataclass(frozen=True)
class GeneratedFile:
    filename: str
    path: Path
    contractnli_id: int | None


def load_manifest() -> Manifest:
    raw = yaml.safe_load(
        resources.files("compliance_seed").joinpath("manifest.yaml").read_text("utf-8")
    )
    vendors = [
        VendorSpec(
            name=v["name"],
            data_access=v["data_access"],
            renewal_date=v["renewal_date"],
            contract_file=v["contract_file"],
            contractnli_id=v["source"]["contractnli_id"] if v["source"] else None,
            source_format=v["source"]["format"] if v["source"] else None,
        )
        for v in raw["vendors"]
    ]
    return Manifest(dataset=raw["dataset"], demo=raw["demo"], vendors=vendors)


def _attribution(doc_id: int, file_name: str) -> str:
    return (
        f"ContractNLI test-split document {doc_id} ({file_name}); "
        "Koreeda & Manning (2021); licensed CC BY 4.0; re-formatted for the demo."
    )


def normalize_zip(data: bytes) -> bytes:
    """Rewrite a zip with fixed timestamps (and fixed Office core-property dates)."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                payload = _CORE_DATE.sub(rb"\g<1>2021-10-05T00:00:00Z\g<2>", payload)
            entry = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            dst.writestr(entry, payload)
    return out.getvalue()


def _pdf_metadata(title: str, subject: str) -> dict[str, str]:
    return {
        "title": title,
        "author": "ContractNLI (Koreeda & Manning, 2021)",
        "subject": subject,
        "keywords": "ContractNLI, CC BY 4.0, demo",
        "creator": "compliance-workspace seed",
        "producer": "compliance-workspace seed",
        "creationDate": FIXED_PDF_DATE,
        "modDate": FIXED_PDF_DATE,
    }


def original_pdf(raw: bytes, title: str, subject: str) -> bytes:
    with pymupdf.open(stream=raw, filetype="pdf") as doc:
        doc.set_metadata(_pdf_metadata(title, subject))
        return bytes(doc.tobytes(garbage=3, deflate=True, no_new_id=True))


def scanned_pdf(raw: bytes, title: str, subject: str, seed: int) -> bytes:
    """Rasterise every page into a slightly skewed, noisy grayscale JPEG scan.

    The result contains only images: no text layer, so ingestion must OCR it.
    """
    rng = np.random.default_rng(seed)
    with pymupdf.open(stream=raw, filetype="pdf") as source, pymupdf.open() as scan:
        for page in source:
            pix = page.get_pixmap(dpi=SCAN_DPI, colorspace=pymupdf.csGRAY, alpha=False)
            image = Image.frombytes("L", (pix.width, pix.height), pix.samples)
            image = image.rotate(0.4, resample=Image.Resampling.BICUBIC, fillcolor=255)
            pixels = np.asarray(image, dtype=np.int16)
            noise = rng.normal(0.0, 9.0, size=pixels.shape)
            noisy = np.clip(pixels + noise, 0, 255).astype(np.uint8)
            buffer = io.BytesIO()
            Image.fromarray(noisy, mode="L").save(buffer, format="JPEG", quality=72)
            target = scan.new_page(width=page.rect.width, height=page.rect.height)
            target.insert_image(target.rect, stream=buffer.getvalue())
        scan.set_metadata(_pdf_metadata(title, subject + " Rasterised to simulate a scanned copy."))
        return bytes(scan.tobytes(garbage=3, deflate=True, no_new_id=True))


_DOCX_HEADING = re.compile(r"^(\d{1,2}(\.\d{1,2})*\.?\s+\S|[A-Z][A-Z\s&,'-]{3,60}$)")


def docx_from_text(text: str, title: str, subject: str) -> bytes:
    document = docx.Document()
    props = document.core_properties
    props.title = title
    props.subject = subject
    props.author = "ContractNLI (Koreeda & Manning, 2021)"
    props.comments = subject
    props.created = FIXED_DATETIME
    props.modified = FIXED_DATETIME
    props.last_modified_by = "compliance-workspace seed"
    props.revision = 1
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    for index, line in enumerate(lines):
        if index == 0:
            document.add_heading(line, level=0)
        elif len(line) <= 80 and _DOCX_HEADING.match(line) and not line.endswith((".", ",", ";")):
            document.add_heading(line, level=2)
        else:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return normalize_zip(buffer.getvalue())


def vendor_register(vendors: list[VendorSpec], sources: dict[int, str]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # noqa: S101 - a new workbook always has an active sheet
    sheet.title = "Vendors"
    sheet["A1"] = "Vendor register - SYNTHETIC demo data generated by make seed"
    sheet["A1"].font = Font(bold=True, size=13)
    headers = ["Vendor name", "Data access level", "Contract file name", "Renewal date"]
    sheet.append([])
    sheet.append(headers)
    for cell in sheet[3]:
        cell.font = Font(bold=True)
    for vendor in vendors:
        sheet.append([vendor.name, vendor.data_access, vendor.contract_file, vendor.renewal_date])
    for row in sheet.iter_rows(min_row=4, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
    for column, width in zip("ABCD", (42, 18, 36, 14), strict=True):
        sheet.column_dimensions[column].width = width
    provenance = workbook.create_sheet("Sources")
    provenance.append(
        ["Contract file name", "ContractNLI test document", "Original file", "License"]
    )
    for vendor in vendors:
        if vendor.contractnli_id is not None:
            provenance.append(
                [
                    vendor.contract_file,
                    vendor.contractnli_id,
                    sources[vendor.contractnli_id],
                    "CC BY 4.0",
                ]
            )
    workbook.properties.created = FIXED_DATETIME
    workbook.properties.modified = FIXED_DATETIME
    workbook.properties.creator = "compliance-workspace seed"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return normalize_zip(buffer.getvalue())


def generate_seed_files(dataset_zip: Path, output_dir: Path) -> list[GeneratedFile]:
    """Write the vendor register and every available contract into ``output_dir``."""
    manifest = load_manifest()
    output_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(dataset_zip) as archive:
        split = json.loads(archive.read(manifest.dataset["split_file"]))
        test_docs = {int(d["id"]): d for d in split["documents"]}
        generated: list[GeneratedFile] = []
        sources: dict[int, str] = {}
        for vendor in manifest.vendors:
            if vendor.contractnli_id is None or vendor.source_format is None:
                continue
            doc = test_docs.get(vendor.contractnli_id)
            if doc is None:
                raise ValueError(
                    f"document {vendor.contractnli_id} is not in the ContractNLI test split"
                )
            sources[vendor.contractnli_id] = doc["file_name"]
            title = f"{vendor.name} - confidentiality agreement"
            subject = _attribution(vendor.contractnli_id, doc["file_name"])
            if vendor.source_format == "docx":
                data = docx_from_text(doc["text"], title, subject)
            else:
                raw = archive.read("contract-nli/raw/" + doc["file_name"])
                if vendor.source_format == "scanned_pdf":
                    data = scanned_pdf(raw, title, subject, seed=vendor.contractnli_id)
                else:
                    data = original_pdf(raw, title, subject)
            path = output_dir / vendor.contract_file
            path.write_bytes(data)
            generated.append(GeneratedFile(vendor.contract_file, path, vendor.contractnli_id))
    register_path = output_dir / manifest.demo["register_filename"]
    register_path.write_bytes(vendor_register(manifest.vendors, sources))
    return [GeneratedFile(register_path.name, register_path, None), *generated]
